# Regenerates the synthetic PPTX used by the live import/reuse scenario.
# Requires python-pptx; generated fixture contains no user content.
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from pathlib import Path
out = Path(__file__).with_name('brand-template.pptx')
p = Presentation()
p.slide_width = Inches(13.333333)
p.slide_height = Inches(7.5)
for title, body in [('Brand introduction', 'Reference placeholder, not factual evidence'), ('Brand details', 'Replace this reference text with authorized source content')]:
    slide = p.slides.add_slide(p.slide_layouts[6])
    for text, top, size in [(title,1,40),(body,3,24)]:
        box = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(11), Inches(1.5))
        para = box.text_frame.paragraphs[0]
        para.text = text
        para.font.name = 'Arial'
        para.font.size = Pt(size)
        para.font.color.rgb = RGBColor.from_string('243B53')
p.save(out)
with ZipFile(out) as z: parts = {name:z.read(name) for name in z.namelist()}
ns = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
theme = etree.fromstring(parts['ppt/theme/theme1.xml'])
theme.set('name','Template reuse acceptance')
for key,color in [('dk1','243B53'),('lt1','FFFFFF'),('accent1','1F7A8C')]:
    node = theme.find('.//a:clrScheme/a:'+key,ns)
    for child in list(node): node.remove(child)
    etree.SubElement(node,'{'+ns['a']+'}srgbClr',val=color)
for key in ['majorFont','minorFont']:
    theme.find('.//a:'+key+'/a:latin',ns).set('typeface','Arial')
parts['ppt/theme/theme1.xml'] = etree.tostring(theme,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(out,'w',ZIP_DEFLATED) as z:
    for name,data in parts.items(): z.writestr(name,data)
print(out.name, out.stat().st_size)
